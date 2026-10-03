"""Offline import tests: no production credentials or customer records required."""
import ast
import base64
import copy
import importlib.util
import io
import json
from pathlib import Path
import sys
import types
import unittest
import zipfile
from datetime import datetime, timezone
from decimal import Decimal

ROOT = Path(__file__).parent


class Conflict(Exception):
    response = {"Error": {"Code": "TransactionCanceledException"}}


class Store:
    def __init__(self):
        self.tables = {}
        self.fail = False
        self.transactions = 0

    def Table(self, name):
        database = self.tables.setdefault(name, {})
        def get_item(Key, **kwargs):
            return {"Item": copy.deepcopy(database[next(iter(Key.values()))])} if next(iter(Key.values())) in database else {}
        def scan(**kwargs):
            rows = list(database.values())
            if kwargs.get('FilterExpression'):
                rows = [r for r in rows if 'FileName' in r]
            return {'Items':copy.deepcopy(rows)}
        return types.SimpleNamespace(get_item=get_item,scan=scan)

    def transact_write_items(self, TransactItems):
        self.transactions += 1
        if self.fail:
            raise Conflict()
        pending = copy.deepcopy(self.tables)
        for operation in TransactItems:
            p = operation['Put']; item=p['Item']; name=p['TableName']
            primary={'companies':'CompanyID','sites':'SiteID','contacts':'ContactID','leads':'LeadID','settings':'SettingKey','owners':'OwnerID','managers':'MgmtCompanyID'}[name]
            db=pending.setdefault(name, {}); old=db.get(item[primary])
            if 'attribute_not_exists' in p['ConditionExpression']:
                if old is not None: raise Conflict()
            else:
                for alias,field in p['ExpressionAttributeNames'].items():
                    if old is None or old.get(field) != p['ExpressionAttributeValues'][':'+alias[1:]]: raise Conflict()
            if len(json.dumps(item, default=str).encode()) > 400000: raise AssertionError('Oversize DynamoDB item')
            db[item[primary]]=copy.deepcopy(item)
        self.tables=pending


class FormsTest(unittest.TestCase):
    def setUp(self):
        self.store=Store()
        api=types.ModuleType('novara_api')
        api.LEADS_TABLE_NAME='leads'; api.SETTINGS_TABLE_NAME='settings'
        api.ensure_leads_table=lambda:None; api.ensure_settings_table=lambda:None
        api.dynamodb_table=self.store.Table
        api.dynamodb_resource=lambda: types.SimpleNamespace(meta=types.SimpleNamespace(client=self.store))
        api.optional_auth_user=lambda headers: ({'role':'aem','userId':'test'},None) if headers else (None,None)
        api.json_safe=lambda x: json.loads(json.dumps(x,default=lambda v: float(v) if isinstance(v,Decimal) else str(v)))
        api._LOGGER=types.SimpleNamespace(exception=lambda *args:None)
        # Exercise the real lead validator, without importing unrelated AWS/pool integrations.
        tree=ast.parse((ROOT/'novara_api.py').read_text())
        names={'parse_lead_payload','LEAD_SOURCES','LEAD_SYSTEM_TYPES','LEAD_STAGES'}
        nodes=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in names or isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id in names for t in n.targets)]
        scope={'Decimal':Decimal,'datetime':datetime,'timezone':timezone,'_as_text':lambda v:'' if v is None else str(v).strip()}
        exec(compile(ast.Module(body=nodes,type_ignores=[]),'real-validator','exec'),scope)
        api.parse_lead_payload=scope['parse_lead_payload']
        self.originals={n:sys.modules.get(n) for n in ('novara_api','master_data_api')}
        sys.modules['novara_api']=api
        spec=importlib.util.spec_from_file_location('master_data_api',ROOT/'master_data_api.py')
        master=importlib.util.module_from_spec(spec); sys.modules['master_data_api']=master;spec.loader.exec_module(master)
        master.COMPANIES_TABLE_NAME='companies';master.SITES_TABLE_NAME='sites';master.CONTACTS_TABLE_NAME='contacts'
        master.ensure_companies_table=lambda:None;master.ensure_contacts_table=lambda:None
        spec=importlib.util.spec_from_file_location('forms_under_test',ROOT/'customer_forms.py')
        self.forms=importlib.util.module_from_spec(spec);spec.loader.exec_module(self.forms)
        self.api=api
        self.payload={'fileName':'Example.csv','fileBase64':base64.b64encode(b'Site name: Example Pool\nCompany name: Example Company').decode(),
            'fields':{'companyName':'Example Company','siteName':'Example Pool','address':'123 Test St','city':'Test City',
                      'contactName':'Test Contact','email':'test@example.com','notes':'Called October 2. Send information.','followUp':'2026-10-05'}}

    def tearDown(self):
        for n,v in self.originals.items():
            if v is None:sys.modules.pop(n,None)
            else:sys.modules[n]=v

    def save(self, payload=None):
        return self.forms.commit(payload or self.payload,{'userId':'test'})

    def test_create_attach_download_retry(self):
        result=self.save()
        for table in ['companies','sites','contacts','leads']:self.assertEqual(len(self.store.tables[table]),1)
        status,download=self.forms.route('GET','/api/customer-forms/download',headers={'x':'y'},query={'id':[result['formId']]})
        self.assertEqual(status,200);self.assertEqual(download['fileBase64'],self.payload['fileBase64'])
        transactions=self.store.transactions
        self.assertTrue(self.save()['alreadyImported']);self.assertEqual(transactions,self.store.transactions)

    def test_preserve_notes_metadata_and_blank_contact(self):
        result=self.save(); lead=self.store.tables['leads'][result['leadId']]
        lead['OwnerID']='OWNER1';lead['EstimatedSavings']=Decimal('1234.50')
        site=self.store.tables['sites'][result['siteId']];site['Systems']=5;site['CustomField']='keep'
        payload=copy.deepcopy(self.payload);payload['fileBase64']=base64.b64encode(b'Updated form').decode()
        payload['fields'].update(notes='Second call',email='',phone='',address='123 Test Street')
        self.save(payload)
        lead=self.store.tables['leads'][result['leadId']]
        self.assertIn('Called October 2',lead['Notes']);self.assertIn('Second call',lead['Notes'])
        self.assertEqual(lead['OwnerID'],'OWNER1');self.assertEqual(lead['EstimatedSavings'],Decimal('1234.50'))
        self.assertEqual(lead['ContactEmail'],'test@example.com')
        site=self.store.tables['sites'][result['siteId']];self.assertEqual(site['Systems'],5);self.assertEqual(site['CustomField'],'keep')

    def test_ambiguous_site_blocks(self):
        self.save();site=copy.deepcopy(next(iter(self.store.tables['sites'].values())))
        site['SiteID']='OTHER';self.store.tables['sites']['OTHER']=site
        with self.assertRaisesRegex(ValueError,'More than one site'):self.save()

    def test_mismatched_company_blocks(self):
        result=self.save();self.payload['fields']['companyName']='Wrong company';self.payload['fields']['siteId']=result['siteId']
        with self.assertRaisesRegex(ValueError,'different company'):self.save()

    def test_bad_followup_and_transaction_failure_leave_no_records(self):
        self.payload['fields']['followUp']='bad-date'
        with self.assertRaisesRegex(ValueError,'NextFollowUp'):self.save()
        self.assertFalse(any(self.store.tables.values()))
        self.payload['fields']['followUp']='2026-10-05';self.store.fail=True
        status,result=self.forms.route('POST','/api/customer-forms/import',headers={'x':'y'},body=self.payload)
        self.assertEqual(status,409);self.assertFalse(any(self.store.tables.values()))

    def test_auth_and_role(self):
        self.assertEqual(self.forms.route('POST','/api/customer-forms/import',body=self.payload)[0],401)
        self.api.optional_auth_user=lambda h:({'role':'owner'},None)
        self.assertEqual(self.forms.route('POST','/api/customer-forms/import',headers={'x':'y'},body=self.payload)[0],403)

    def test_docx_xlsx_pdf_csv_extraction(self):
        result=self.forms.preview(self.payload);self.assertEqual(result['fields']['siteName'],'Example Pool')
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w') as z:
            z.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Site name: Example Pool</w:t></w:r></w:p></w:body></w:document>')
        self.assertIn('Example Pool',self.forms.extract('docx',stream.getvalue()))
        from openpyxl import Workbook
        stream=io.BytesIO();book=Workbook();book.active.append(['Site','Pool area']);book.active.append(['Example Pool',950]);book.save(stream)
        self.assertIn('Example Pool | 950',self.forms.extract('xlsx',stream.getvalue()))
        from pypdf import PdfWriter
        stream=io.BytesIO();pdf=PdfWriter();pdf.add_blank_page(width=100,height=100);pdf.write(stream)
        self.assertEqual(self.forms.extract('pdf',stream.getvalue()),'')

    def test_size_limit_and_chunked_original(self):
        self.payload['fileBase64']=base64.b64encode(b'x'*(2*1024*1024)).decode()
        result=self.save();meta=self.store.tables['settings'][result['formId']]
        self.assertGreater(meta['Chunks'],1)
        self.payload['fileBase64']=base64.b64encode(b'x'*(2*1024*1024+1)).decode()
        with self.assertRaisesRegex(ValueError,'2 MB'):self.save()


if __name__=='__main__':unittest.main()
